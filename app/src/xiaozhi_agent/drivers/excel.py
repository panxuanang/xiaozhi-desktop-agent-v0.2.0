from __future__ import annotations

import operator
import re
import shutil
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import PatternFill

from ..paths import OUTBOX_DIR
from .files import resolve_path


RED_FILL = PatternFill(fill_type="solid", fgColor="FFC7CE")


def _sheet(wb, sheet: str | None):
    if sheet and sheet in wb.sheetnames:
        return wb[sheet]
    return wb[wb.sheetnames[0]]


def _out(src: Path, tag: str) -> Path:
    OUTBOX_DIR.mkdir(parents=True, exist_ok=True)
    base = re.sub(r"[^0-9A-Za-z_\-\u4e00-\u9fff]", "_", src.stem)[:80]
    return OUTBOX_DIR / f"{base}_{tag}{src.suffix}"


def _header_map(ws) -> dict[str, int]:
    return {str(c.value).strip(): c.column for c in ws[1] if c.value is not None}


def _col_index(ws, column: str | int) -> int:
    if isinstance(column, int):
        return column
    raw = str(column).strip()
    if raw.isdigit():
        return int(raw)
    headers = _header_map(ws)
    if raw in headers:
        return headers[raw]
    from openpyxl.utils import column_index_from_string
    return column_index_from_string(raw.upper())


def excel_sort(path: str, workspace: Path, sheet: str | None, column: str | int, descending: bool = False) -> Path:
    src = resolve_path(path, workspace)
    wb = load_workbook(src)
    ws = _sheet(wb, sheet)
    ci = _col_index(ws, column)
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    rows.sort(key=lambda r: (r[ci - 1] is None, r[ci - 1]), reverse=bool(descending))
    for r_idx, row in enumerate(rows, start=2):
        for c_idx, value in enumerate(row, start=1):
            ws.cell(r_idx, c_idx).value = value
    out = _out(src, "sorted")
    wb.save(out)
    return out


def excel_dedupe(path: str, workspace: Path, sheet: str | None, columns: list[str | int] | None = None) -> Path:
    src = resolve_path(path, workspace)
    wb = load_workbook(src)
    ws = _sheet(wb, sheet)
    indexes = [_col_index(ws, c) for c in columns] if columns else list(range(1, ws.max_column + 1))
    seen = set()
    delete_rows: list[int] = []
    for r in range(2, ws.max_row + 1):
        key = tuple(ws.cell(r, c).value for c in indexes)
        if key in seen:
            delete_rows.append(r)
        else:
            seen.add(key)
    for r in reversed(delete_rows):
        ws.delete_rows(r, 1)
    out = _out(src, "deduped")
    wb.save(out)
    return out


def excel_replace(path: str, workspace: Path, sheet: str | None, find: Any, replace: Any) -> Path:
    src = resolve_path(path, workspace)
    wb = load_workbook(src)
    ws = _sheet(wb, sheet)
    for row in ws.iter_rows():
        for cell in row:
            if cell.value == find:
                cell.value = replace
            elif isinstance(cell.value, str) and isinstance(find, str) and find in cell.value:
                cell.value = cell.value.replace(find, str(replace))
    out = _out(src, "replaced")
    wb.save(out)
    return out


def excel_fill_blank(path: str, workspace: Path, sheet: str | None, column: str | int, value: Any) -> Path:
    src = resolve_path(path, workspace)
    wb = load_workbook(src)
    ws = _sheet(wb, sheet)
    ci = _col_index(ws, column)
    for r in range(2, ws.max_row + 1):
        if ws.cell(r, ci).value in (None, ""):
            ws.cell(r, ci).value = value
    out = _out(src, "filled")
    wb.save(out)
    return out


def excel_sum(path: str, workspace: Path, sheet: str | None, column: str | int, output_cell: str = "") -> Path:
    src = resolve_path(path, workspace)
    wb = load_workbook(src)
    ws = _sheet(wb, sheet)
    ci = _col_index(ws, column)
    total = 0.0
    for r in range(2, ws.max_row + 1):
        v = ws.cell(r, ci).value
        if isinstance(v, (int, float)):
            total += float(v)
    if output_cell:
        ws[output_cell] = total
    else:
        ws.cell(ws.max_row + 2, ci).value = total
    out = _out(src, "sum")
    wb.save(out)
    return out


def excel_compute_column(path: str, workspace: Path, sheet: str | None, target_column: str | int, formula_template: str) -> Path:
    src = resolve_path(path, workspace)
    wb = load_workbook(src)
    ws = _sheet(wb, sheet)
    ci = _col_index(ws, target_column)
    for r in range(2, ws.max_row + 1):
        ws.cell(r, ci).value = str(formula_template).replace("{row}", str(r))
    out = _out(src, "computed")
    wb.save(out)
    return out


def excel_conditional_red(path: str, workspace: Path, sheet: str | None, column: str | int, operator_name: str, value: Any) -> Path:
    src = resolve_path(path, workspace)
    wb = load_workbook(src)
    ws = _sheet(wb, sheet)
    ci = _col_index(ws, column)
    ops = {">": operator.gt, ">=": operator.ge, "<": operator.lt, "<=": operator.le, "==": operator.eq, "!=": operator.ne}
    op = ops.get(operator_name)
    if not op:
        raise ValueError(f"不支持的比较符: {operator_name}")
    for r in range(2, ws.max_row + 1):
        cell = ws.cell(r, ci)
        try:
            if op(cell.value, value):
                cell.fill = RED_FILL
        except TypeError:
            continue
    out = _out(src, "marked")
    wb.save(out)
    return out


def excel_merge_files(paths: list[str], workspace: Path, output: str = "merged.xlsx") -> Path:
    resolved = [resolve_path(p, workspace) for p in paths]
    if not resolved:
        raise ValueError("没有输入文件")
    out = resolve_path(output, OUTBOX_DIR) if Path(output).is_absolute() else OUTBOX_DIR / output
    OUTBOX_DIR.mkdir(parents=True, exist_ok=True)
    wb_out = Workbook()
    ws_out = wb_out.active
    ws_out.title = "Merged"
    wrote_header = False
    for src in resolved:
        wb = load_workbook(src, data_only=False)
        ws = wb[wb.sheetnames[0]]
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            continue
        if not wrote_header:
            ws_out.append(list(rows[0]))
            wrote_header = True
        for row in rows[1:]:
            ws_out.append(list(row))
    wb_out.save(out)
    return out


def excel_lookup(
    left_path: str,
    right_path: str,
    workspace: Path,
    left_key: str,
    right_key: str,
    right_value: str,
    output: str = "lookup.xlsx",
) -> Path:
    left = resolve_path(left_path, workspace)
    right = resolve_path(right_path, workspace)
    wb_l = load_workbook(left)
    ws_l = wb_l[wb_l.sheetnames[0]]
    wb_r = load_workbook(right, data_only=True)
    ws_r = wb_r[wb_r.sheetnames[0]]
    hl = _header_map(ws_l)
    hr = _header_map(ws_r)
    li, rki, rvi = hl[left_key], hr[right_key], hr[right_value]
    mapping = {ws_r.cell(r, rki).value: ws_r.cell(r, rvi).value for r in range(2, ws_r.max_row + 1)}
    target_col = ws_l.max_column + 1
    ws_l.cell(1, target_col).value = right_value
    for r in range(2, ws_l.max_row + 1):
        ws_l.cell(r, target_col).value = mapping.get(ws_l.cell(r, li).value)
    OUTBOX_DIR.mkdir(parents=True, exist_ok=True)
    out = OUTBOX_DIR / Path(output).name
    wb_l.save(out)
    return out


def excel_split_by_column(path: str, workspace: Path, sheet: str | None, column: str | int, output_dir: str = "") -> list[Path]:
    src = resolve_path(path, workspace)
    wb = load_workbook(src, data_only=False)
    ws = _sheet(wb, sheet)
    ci = _col_index(ws, column)
    groups: dict[str, list[tuple[Any, ...]]] = {}
    header = tuple(c.value for c in ws[1])
    for row in ws.iter_rows(min_row=2, values_only=True):
        key = str(row[ci - 1] if row[ci - 1] is not None else "blank")
        groups.setdefault(key, []).append(tuple(row))
    outdir = resolve_path(output_dir, workspace) if output_dir else OUTBOX_DIR / f"{src.stem}_split"
    outdir.mkdir(parents=True, exist_ok=True)
    outputs: list[Path] = []
    for key, rows in groups.items():
        safe = re.sub(r'[<>:"/\\|?*\x00-\x1f]', "_", key)[:80] or "blank"
        owb = Workbook()
        ows = owb.active
        ows.append(header)
        for row in rows:
            ows.append(row)
        out = outdir / f"{safe}.xlsx"
        owb.save(out)
        outputs.append(out)
    return outputs


def excel_split_sheets(path: str, workspace: Path, output_dir: str = "") -> list[Path]:
    src = resolve_path(path, workspace)
    wb = load_workbook(src, data_only=False)
    outdir = resolve_path(output_dir, workspace) if output_dir else OUTBOX_DIR / f"{src.stem}_sheets"
    outdir.mkdir(parents=True, exist_ok=True)
    outputs: list[Path] = []
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        owb = Workbook()
        ows = owb.active
        ows.title = sheet_name[:31]
        for row in ws.iter_rows(values_only=False):
            ows.append([c.value for c in row])
        safe = re.sub(r'[<>:"/\\|?*\x00-\x1f]', "_", sheet_name)[:80]
        out = outdir / f"{safe}.xlsx"
        owb.save(out)
        outputs.append(out)
    return outputs


def _compare(lhs: Any, operator_name: str, rhs: Any) -> bool:
    ops = {">": operator.gt, ">=": operator.ge, "<": operator.lt, "<=": operator.le, "==": operator.eq, "!=": operator.ne}
    op = ops.get(str(operator_name))
    if not op:
        raise ValueError(f"不支持的比较符: {operator_name}")
    try:
        return bool(op(lhs, rhs))
    except TypeError:
        return False


def excel_filter(path: str, workspace: Path, sheet: str | None, column: str | int, operator: str, value: Any) -> Path:
    src = resolve_path(path, workspace)
    wb = load_workbook(src, data_only=False)
    ws = _sheet(wb, sheet)
    ci = _col_index(ws, column)
    header = [c.value for c in ws[1]]
    rows = [tuple(c.value for c in row) for row in ws.iter_rows(min_row=2) if _compare(row[ci-1].value, operator, value)]
    owb = Workbook(); ows = owb.active; ows.title = ws.title[:31]
    ows.append(header)
    for row in rows: ows.append(row)
    out = _out(src, "filtered")
    owb.save(out)
    return out


def excel_conditional_set(path: str, workspace: Path, sheet: str | None, condition_column: str | int, operator: str, value: Any, target_column: str | int, target_value: Any) -> Path:
    src = resolve_path(path, workspace)
    wb = load_workbook(src)
    ws = _sheet(wb, sheet)
    cci = _col_index(ws, condition_column)
    tci = _col_index(ws, target_column)
    for r in range(2, ws.max_row + 1):
        if _compare(ws.cell(r, cci).value, operator, value):
            ws.cell(r, tci).value = target_value
    out = _out(src, "conditional_set")
    wb.save(out)
    return out


def excel_batch_replace(paths: list[str], workspace: Path, find: Any, replace: Any) -> list[Path]:
    outputs: list[Path] = []
    for path in paths:
        outputs.append(excel_replace(path, workspace, None, find, replace))
    return outputs
