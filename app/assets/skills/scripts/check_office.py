from __future__ import annotations

import sys
from pathlib import Path


def check(path: Path) -> None:
    ext = path.suffix.lower()
    if ext == '.xlsx':
        from openpyxl import load_workbook
        wb = load_workbook(path, read_only=True, data_only=False)
        print(f'OK XLSX sheets={len(wb.sheetnames)}')
        wb.close()
    elif ext == '.docx':
        from docx import Document
        doc = Document(path)
        print(f'OK DOCX paragraphs={len(doc.paragraphs)} tables={len(doc.tables)}')
    elif ext == '.pptx':
        from pptx import Presentation
        prs = Presentation(path)
        print(f'OK PPTX slides={len(prs.slides)}')
    else:
        print(f'SKIP {path.name}')


if __name__ == '__main__':
    for arg in sys.argv[1:]:
        check(Path(arg))
